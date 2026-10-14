"""Agrupa la población municipal por sexo y edad para 2007–2025."""

from __future__ import annotations

import argparse
import csv
import io
import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = ROOT / "data/raw/population"
DEFAULT_OUTPUT = ROOT / "data/processed/population/population_municipal_2007_2025.csv"
LOGGER = logging.getLogger(__name__)
IDENTIFIERS = ("DP", "DPNOM", "DPMP", "MPIO", "AÑO", "ÁREA GEOGRÁFICA")
SEXES = ("Hombres", "Mujeres", "Total")
AGE_GROUPS = {
    "<1": range(0, 1),
    "1-4": range(1, 5),
    "5-14": range(5, 15),
    "15-49": range(15, 50),
}
SOURCES = (
    ("DCD-area-sexo-edad-proypoblacion-Mun-2005-2017_VP.xlsx", "NuevaMpal", 12),
    ("PPED-AreaSexoEdadMun-2018-2042_VP.xlsx", "PobMunicipalxÁreaSexoEdad", 9),
)
OUTPUT_COLUMNS = (*IDENTIFIERS, *(
    f"{sex}_{group}" for sex in SEXES for group in (*AGE_GROUPS, "50+")
))


def source_columns(header: list[str], recent: bool) -> dict[str, list[int]]:
    """Resuelve una vez las columnas de cada grupo de edad."""
    positions = {name: index for index, name in enumerate(header)}
    for name in IDENTIFIERS:
        if name not in positions:
            raise ValueError(f"Falta la columna {name!r}")

    grouped = {}
    for sex in SEXES:
        for group, ages in (*AGE_GROUPS.items(), ("50+", range(50, 100 if recent else 85))):
            if recent:
                names = [f"{sex} {age} {'año' if age == 1 else 'años'}" for age in ages]
            else:
                names = [f"{sex}_{age}" for age in ages]
            if group == "50+":
                if recent:
                    open_ended = [name for name in header
                                  if name.startswith(f"{sex} 100 años")
                                  and name.endswith("y más")]
                    if len(open_ended) != 1:
                        raise ValueError(f"Se esperaba una columna {sex} 100 años y más")
                    names.extend(open_ended)
                else:
                    names.append(f"{sex}_85 y más")
            missing = [name for name in names if name not in positions]
            if missing:
                raise ValueError(f"Faltan columnas de edad: {missing}")
            grouped[f"{sex}_{group}"] = [positions[name] for name in names]
    return grouped


def write_source(source: Path, sheet_name: str, header_row: int,
                 writer: csv.writer) -> int:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = workbook[sheet_name].iter_rows(values_only=True)
        header = None
        for number, row in enumerate(rows, start=1):
            if number == header_row - 1 and header_row == 9:
                first_header = row
            if number == header_row:
                if header_row == 9:
                    header = list(first_header[:6]) + list(row[6:])
                else:
                    header = list(row)
                break
        if header is None:
            raise ValueError(f"{source.name}: no se encontró el encabezado")

        positions = {name: index for index, name in enumerate(header)}
        groups = source_columns(header, recent=header_row == 9)
        count = 0
        for row in rows:
            year = row[positions["AÑO"]]
            if not isinstance(year, (int, float)) or not 2007 <= year <= 2025:
                continue

            identifiers = [row[positions[name]] for name in IDENTIFIERS]
            code = identifiers[3]
            identifiers[3] = str(int(code) if isinstance(code, (int, float)) else code).zfill(5)
            identifiers[4] = int(year)

            values = []
            for indices in groups.values():
                numbers = [row[index] for index in indices if row[index] is not None]
                values.append(sum(numbers) if numbers else None)
            writer.writerow((*identifiers, *values))
            count += 1
        return count
    finally:
        workbook.close()


def aggregate_for_join(groups: pd.DataFrame) -> pd.DataFrame:
    """Reproduce poblacion_to_join del notebook: una fila por municipio y año."""
    df = groups.copy()
    df["codigo_municipio"] = df["MPIO"].astype("string").str.strip().str.zfill(5)
    df["ano_epidemiologico"] = pd.to_numeric(df["AÑO"], errors="coerce").astype("Int64")

    age_groups = ("<1", "1-4", "5-14", "15-49", "50+")
    for sex in SEXES:
        for age_group in age_groups:
            name = f"{sex}_{age_group}"
            df[name] = pd.to_numeric(df[name], errors="coerce")

    is_total = df["ÁREA GEOGRÁFICA"].eq("Total")
    is_urban = df["ÁREA GEOGRÁFICA"].eq("Cabecera Municipal")
    is_rural = df["ÁREA GEOGRÁFICA"].eq("Centros Poblados y Rural Disperso")
    result = df[["codigo_municipio", "ano_epidemiologico"]].copy()

    for sex, output_name in (("Total", "poblacion_total"),
                             ("Hombres", "poblacion_hombres"),
                             ("Mujeres", "poblacion_mujeres")):
        columns = [f"{sex}_{age_group}" for age_group in age_groups]
        result[output_name] = df[columns].sum(axis=1, min_count=5).where(is_total)

    total_columns = [f"Total_{age_group}" for age_group in age_groups]
    total_by_area = df[total_columns].sum(axis=1, min_count=5)
    result["poblacion_cabecera"] = total_by_area.where(is_urban)
    result["poblacion_rural"] = total_by_area.where(is_rural)

    age_suffixes = ("menor_1", "1_4", "5_14", "15_49", "50_mas")
    for sex, prefix in (("Total", "poblacion"),
                        ("Hombres", "poblacion_hombres"),
                        ("Mujeres", "poblacion_mujeres")):
        for age_group, suffix in zip(age_groups, age_suffixes):
            result[f"{prefix}_{suffix}"] = df[f"{sex}_{age_group}"].where(is_total)

    value_columns = list(result.columns[2:])
    aggregated = result.groupby(
        ["codigo_municipio", "ano_epidemiologico"], as_index=False
    )[value_columns].max()
    aggregated[value_columns] = aggregated[value_columns].astype("Int64")
    return aggregated


def process(input_dir: Path, output: Path) -> None:
    # Las filas agrupadas son intermedias: solo se escribe la tabla municipal final.
    intermediate = io.StringIO()
    writer = csv.writer(intermediate)
    writer.writerow(OUTPUT_COLUMNS)
    for filename, sheet_name, header_row in SOURCES:
        source = input_dir / filename
        LOGGER.info("Leyendo %s", source)
        count = write_source(source, sheet_name, header_row, writer)
        LOGGER.info("%s: %d filas (2007–2025)", filename, count)
    intermediate.seek(0)
    groups = pd.read_csv(intermediate, dtype={"DP": "string", "MPIO": "string"})
    joined = aggregate_for_join(groups)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.name + ".part")
    try:
        joined.to_csv(temporary, index=False, encoding="utf-8")
        temporary.replace(output)
        LOGGER.info("Guardado %s: %d municipios-año", output, len(joined))
    finally:
        temporary.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    for filename, _, _ in SOURCES:
        if not (args.input_dir / filename).is_file():
            parser.error(f"No existe {args.input_dir / filename}")
    process(args.input_dir, args.output)


if __name__ == "__main__":
    main()
