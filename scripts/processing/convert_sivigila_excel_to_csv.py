"""Convierte los libros Excel de Sivigila en archivos CSV, uno por año."""

from __future__ import annotations

import argparse
import csv
import logging
from pathlib import Path

from openpyxl import load_workbook


LOGGER = logging.getLogger(__name__)
ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = ROOT / "data/raw/sivigila/excel"
DEFAULT_OUTPUT = ROOT / "data/raw/sivigila/csv"
SHEET_NAME = "Datos"


def convert_file(source: Path, output_dir: Path, overwrite: bool = False) -> Path:
    """Escribe la hoja Datos en UTF-8, sin cargar el libro completo en memoria."""
    destination = output_dir / f"{source.stem}.csv"
    if destination.exists() and not overwrite:
        LOGGER.info("Ya existe; se omite: %s", destination)
        return destination

    output_dir.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(destination.name + ".part")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != [SHEET_NAME]:
            raise ValueError(
                f"{source.name}: se esperaba una sola hoja llamada {SHEET_NAME!r}; "
                f"hay {workbook.sheetnames}"
            )
        sheet = workbook[SHEET_NAME]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None or not any(value is not None for value in header):
            raise ValueError(f"{source.name}: la hoja no tiene encabezado")
        count = 0
        with temporary.open("w", encoding="utf-8", newline="") as output:
            writer = csv.writer(output)
            writer.writerow(header)
            for row in rows:
                if any(value is not None for value in row):
                    writer.writerow(row)
                    count += 1
        temporary.replace(destination)
        LOGGER.info("Guardado %s: %d registros, %d columnas",
                    destination, count, len(header))
        return destination
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--overwrite", action="store_true",
                        help="Regenerar los CSV existentes")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")

    if not args.input_dir.is_dir():
        parser.error(f"No existe la carpeta de Excel: {args.input_dir}")
    sources = sorted(path for path in args.input_dir.glob("*.xlsx")
                     if not path.name.startswith("~$"))
    if not sources:
        parser.error(f"No se encontraron archivos .xlsx en {args.input_dir}")
    LOGGER.info("Convirtiendo %d archivos de %s a %s",
                len(sources), args.input_dir, args.output_dir)
    for source in sources:
        convert_file(source, args.output_dir, args.overwrite)


if __name__ == "__main__":
    main()
